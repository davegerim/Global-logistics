"""Generate Intellect LMS v22 Study Guide as DOCX."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_shading(cell, color_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(11)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    doc.add_paragraph(text, style=style)


def add_table(doc, headers, rows, header_color="1F4E79"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], header_color)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()
    return table


def add_code_block(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def build_document():
    doc = Document()

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Intellect Design Arena\nLMS v22 Study Guide")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(31, 78, 121)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(
        "eMACH.ai Lending Ecosystem\n"
        "(LOS · LMS · CLMS · Collections · Integration)"
    )
    sr.font.size = Pt(14)
    sr.italic = True

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        "\n\nPrepared for Intellect Design Arena Training\n"
        "Based on public Intellect documentation + standard LMS architecture\n"
        "August 2026"
    )
    mr.font.size = Pt(11)
    mr.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_page_break()

    # Important notice
    add_heading(doc, "Important Notice", 1)
    add_bullet(doc, "LMS = Loan Management System (post-disbursement servicing), NOT a learning/training LMS.")
    add_bullet(doc, "Version 22 is your company's internal release label. Server URLs, credentials, and v22-only screen differences must come from your trainer or internal wiki.")
    add_bullet(doc, "Do NOT share internal server IPs, passwords, or client bank data outside Intellect.")
    add_bullet(doc, "Add your own training screenshots under Section 11 (Screenshot Checklist).")

    doc.add_page_break()

    # TOC
    add_heading(doc, "Table of Contents", 1)
    toc_items = [
        "1. What is LMS v22?",
        "2. Platform Architecture",
        "3. All Related Modules (LOS, CLMS, etc.)",
        "4. End-to-End Credit Lifecycle Flow",
        "5. LMS Internal Structure",
        "6. Screen & Function Codes (G / B / K codes)",
        "7. Product Configuration (Product Engine)",
        "8. Daily Transactions vs EOD Batch",
        "9. Integration Layer (iTurmeric / Olive Fabric)",
        "10. Server & Environment Access",
        "11. Study Checklist & Screenshots",
        "12. Glossary",
        "13. Official References",
    ]
    for item in toc_items:
        add_bullet(doc, item)

    doc.add_page_break()

    # Section 1
    add_heading(doc, "1. What is LMS v22?", 1)
    add_table(doc, ["Term", "Meaning at Intellect"], [
        ["LMS", "Loan Management System — post-disbursement servicing"],
        ["v22", "Internal product release (e.g. 22.x). Marketing name: eMACH.ai Retail/Commercial Loan Management"],
        ["Legacy name", "Earlier deployments: Polaris Intellect LMS"],
        ["Business unit", "iGCB (Intellect Global Consumer Banking) for retail"],
    ])
    add_para(doc, "LMS owns: booking → disbursement → interest → dues → repayment → delinquency → accounting → closure.")
    add_para(doc, "LMS does NOT own: loan application and credit committee approval (that is LOS).")

    # Section 2
    add_heading(doc, "2. Platform Architecture", 1)
    add_para(doc, "eMACH.ai Design Pillars:", bold=True)
    add_table(doc, ["Pillar", "What it means for LMS"], [
        ["Microservices", "Booking, interest, EOD, delinquency as separate deployable services"],
        ["API-first", "1200+ open APIs; LMS integrates via APIs not hard-coded links"],
        ["Cloud-native", "On-prem, private cloud, public cloud (AWS/Azure/GCP), hybrid"],
        ["Headless", "UI (DEP/CBX-R) separate from LMS logic"],
        ["Event-driven", "Loan booked, payment received, DPD crossed → events to other systems"],
        ["Embedded AI", "PF Credit Digital Experts on Purple Fabric"],
    ])

    add_para(doc, "Architecture Diagram (Text):", bold=True)
    add_code_block(doc, """CUSTOMER (Mobile / Web / Branch)
        │
        ▼
   DEP / CBX-R  (Digital journeys: B21–B2C)
        │
        ▼
   LOS (G22) ──► CLMS (G23, G24)
        │
        ▼ (Approved loan handoff)
   LMS v22 (G25 Booking, G26 Servicing, G28 EOD)
        │
        ▼
   iTurmeric / Olive Fabric (APIs, Events, Flows)
        │
        ▼
   Core Banking · GL · Bureau · KYC · CRM · DWH""")

    add_para(doc, "DIY / Customization Stack:", bold=True)
    add_table(doc, ["Tool", "Purpose"], [
        ["Product Engine", "Clone products, masters, parameters"],
        ["Business Rules Engine (DMN)", "Policies, scorecards, delinquency rules"],
        ["Business Central (BPMN 2.0)", "Credit workflows, approvals"],
        ["Credit Customization Studio", "Screens, domain objects, events, alerts"],
        ["iTurmeric Studio", "New APIs, orchestration, Fluid Pages (UI)"],
        ["Reporting", "Standard reports, dashboards, extracts"],
    ])

    # Section 3
    add_heading(doc, "3. All Related Modules", 1)
    add_table(doc, ["Stage", "Module", "Product", "Screen Codes"], [
        ["10", "Application initiation", "LOS", "G22"],
        ["20", "Credit decisioning", "LOS + Rules + AI", "—"],
        ["30", "Approval & monitoring", "LOS + CLMS", "G23, G24"],
        ["50", "Product management", "Product Engine", "G21"],
        ["Booking & servicing", "LMS v22", "Loan Management", "G25, G26"],
        ["70", "Collateral & limits", "CLMS", "G23, G24"],
        ["80", "Debt / collections", "Debt Management", "G27 + Collections"],
        ["Digital", "DEP / CBX-R", "Digital Engagement", "B21–B2C"],
        ["Batch", "LMS batch", "Batch Processes", "G28"],
        ["Reports", "Reporting", "Analytics", "K211–K21C"],
        ["Interfaces", "Integration", "iTurmeric", "K221–K226"],
    ])

    add_heading(doc, "3.1 LOS — Loan Origination System", 2)
    add_para(doc, "Purpose: Apply → verify → score → approve → prepare for booking.")
    add_bullet(doc, "Omni-channel sourcing (self, RM, broker, API)")
    add_bullet(doc, "Data capture: customer, loan, collateral, PD notes")
    add_bullet(doc, "Verification, assessment, policy, scorecard, delegation matrix")
    add_bullet(doc, "Workflow: parallel / sequential / STP, SLA/TAT")
    add_bullet(doc, "Handoff to LMS: sanitized, validated loan data on approval")

    add_heading(doc, "3.2 LMS v22 — Loan Management System", 2)
    add_code_block(doc, """LMS v22
├── Product Creation (setup/masters)
├── Transactions (day-to-day ops)
├── EOD / Batch (automated daily)
├── Booking/Initiation (G25)
├── Maintenance
└── Queries & Reports""")

    add_heading(doc, "3.3 CLMS — Collateral & Limit Management", 2)
    add_para(doc, "Collateral: property, vehicle, equipment, inventory, securities, receivables, guarantees.")
    add_para(doc, "Lifecycle: earmark, freeze/defreeze, pooling, swap, revaluation, release.")
    add_para(doc, "Limits: hierarchical (group → entity → product → facility), sanctioned vs drawing limits.")

    add_heading(doc, "3.4 Debt Management / Collections", 2)
    add_para(doc, "Flow: Data upload → segmentation → worklist → collector assignment → settlement / legal / repossession.")

    add_heading(doc, "3.5 DEP / CBX-R (Digital Front-end)", 2)
    add_table(doc, ["Code", "Screen"], [
        ["B21", "Portfolio summary"], ["B22", "Account details"], ["B23", "Account statement"],
        ["B24", "Loan repayment"], ["B25", "Recent transactions"], ["B26", "Interest certificate"],
        ["B27", "Disbursement schedule"], ["B28", "New loan request"], ["B29", "Disbursement request"],
        ["B2A", "Reschedule request"], ["B2B", "Top-up request"], ["B2C", "Closure simulation"],
    ])

    # Section 4
    add_heading(doc, "4. End-to-End Credit Lifecycle Flow", 1)
    add_heading(doc, "4.1 Happy Path — New Retail Loan", 2)
    add_bullet(doc, "Customer: B28 New Loan Request via DEP/CBX-R")
    add_bullet(doc, "LOS: KYC, bureau, score, policy checks")
    add_bullet(doc, "CLMS: Limits & collateral validation")
    add_bullet(doc, "LOS: Credit approval")
    add_bullet(doc, "LMS G25: Retail Loan Booking (handoff from LOS)")
    add_bullet(doc, "LMS: Disbursal (single/multi-tranche), repayment schedule generation")
    add_bullet(doc, "Monthly: B24 Loan Repayment → Scheduled Payment → priority allocation → GL")
    add_bullet(doc, "Nightly G28 EOD: accrual, DPD stamping, GL postings")

    add_heading(doc, "4.2 Stage Numbering", 2)
    add_table(doc, ["Stage", "Name", "Primary System"], [
        ["10", "Application initiation", "LOS"],
        ["20", "Credit decisioning", "LOS"],
        ["30", "Approval & monitoring", "LOS + CLMS"],
        ["40", "Advance credit analysis", "LOS"],
        ["50", "Product management", "Product Engine"],
        ["60", "Field operations", "Mobile/LOS"],
        ["70", "Collateral & limits", "CLMS"],
        ["80", "Debt management", "Collections"],
    ])

    add_heading(doc, "4.3 LOS → LMS Handoff (Critical)", 2)
    add_bullet(doc, "Sanitized payload sent on approval (customer, product, amount, tenor, rate, collateral, schedule)")
    add_bullet(doc, "LMS performs Retail/Corporate Loan Booking (G25)")
    add_bullet(doc, "No duplicate data entry in operations")
    add_bullet(doc, "Failure point: approved in LOS but not booked in LMS — ops must re-trigger or manual booking")

    # Section 5
    add_heading(doc, "5. LMS Internal Structure", 1)

    add_heading(doc, "5.1 Product Creation", 2)
    add_table(doc, ["Component", "What you configure"], [
        ["Product Definition", "Loan category (mortgage, PL, auto, WC)"],
        ["Rate Type Scheme", "Fixed, floating, tier, hybrid, adjustable"],
        ["Due knock-off priority", "Tax → fees → penal → interest → principal"],
        ["Fee Scheme", "Processing, prepayment, late fees"],
        ["Credit Scheme", "NPA/classification rules"],
        ["Penal Scheme", "Penalty on overdue"],
        ["Delinquency Scheme", "DPD buckets, movements"],
        ["Accounting Scheme", "GL mapping, accrual rules"],
    ])

    add_heading(doc, "5.2 Booking / Initiation (G25)", 2)
    add_table(doc, ["Function", "Description"], [
        ["Entity Definition", "Legal entity / branch structure"],
        ["Project Definition", "Project finance linkage"],
        ["Retail Loan Booking", "Book approved retail loan"],
        ["Corporate Loan Booking", "Book SME/corporate facility"],
        ["Top-up loan", "Additional on existing loan"],
        ["Rollover loan", "Renew maturing facility"],
        ["Disbursal", "Release funds (single/multi beneficiary)"],
        ["Document creation", "Loan agreement, advice"],
        ["Notepad", "Ops notes on account"],
        ["Payout", "Fund transfer to customer/vendor"],
    ])

    add_heading(doc, "5.3 Transactions (G26)", 2)
    add_table(doc, ["Transaction", "When used"], [
        ["Scheduled Payment", "EMI / installment receipt"],
        ["Payment Reversal", "Reverse wrong posting"],
        ["Excess Adjustment", "Apply surplus to principal or refund"],
        ["Adhoc Fees & Waivers", "One-off charge or waiver"],
        ["Manual Delinquency", "Force/correct DPD bucket"],
        ["Closure (Maturity & Early)", "Normal or prepayment closure"],
        ["Loan & Disbursal Rescheduling", "Change tenor/EMI/disbursement plan"],
        ["Deferral", "Skip/postpone installment (moratorium)"],
        ["Waivers", "Interest/fee waiver"],
    ])

    add_heading(doc, "5.4 Maintenance", 2)
    add_table(doc, ["Function", "Purpose"], [
        ["Associate Details", "Co-borrower/guarantor updates"],
        ["Collateral Details", "Link/update security"],
        ["Disbursal Hold", "Block further disbursement"],
        ["Interest Rate Change", "Manual or benchmark reset"],
        ["Link-Delink Fees", "Attach/detach fee types"],
        ["POS Change", "Principal outstanding adjustment"],
        ["Repayment Mode", "NACH, standing instruction, cash"],
        ["Holiday Master", "Non-working days for due dates"],
        ["Bulk Uploads", "Mass payments, deferrals"],
    ])

    add_heading(doc, "5.5 EOD / Batch (G28)", 2)
    add_table(doc, ["Batch Job", "What it does"], [
        ["Interest calculation", "Compute interest for the day/period"],
        ["Accrual – Due creation", "Create installment dues"],
        ["Late payment fees", "Apply penal charges on overdue"],
        ["Delinquency movement", "Move accounts across DPD buckets"],
        ["Auto Closure", "Close zero-balance/matured loans"],
        ["Auto Reschedule", "Apply predefined reschedule rules"],
        ["Provision calculation", "Regulatory provisioning (NPA)"],
        ["Financial postings", "Push entries to GL/sub-ledger"],
    ])

    add_heading(doc, "5.6 Delinquency (G27)", 2)
    add_table(doc, ["Concept", "Meaning"], [
        ["DPD", "Days Past Due — days after due date unpaid"],
        ["MPD", "Months Past Due (some products)"],
        ["Accrual stopping", "Stop recognizing income on NPA"],
        ["Income reversal", "Reverse accrued interest on classification"],
        ["Provisioning", "Set aside amount per regulator rules"],
    ])

    # Section 6
    add_heading(doc, "6. Screen & Function Codes", 1)

    add_heading(doc, "6.1 G-codes — Back-office Lending", 2)
    add_table(doc, ["Code", "Module"], [
        ["G21", "Products / Product Engine"],
        ["G22", "Origination (LOS)"],
        ["G23", "Customer Limits"],
        ["G24", "Collateral Management"],
        ["G25", "Loan Initiation / Booking"],
        ["G26", "Loan Servicing (core LMS transactions)"],
        ["G27", "Delinquency Management"],
        ["G28", "Batch Processes (EOD)"],
    ])

    add_heading(doc, "6.2 B-codes — Digital Engagement (DEP)", 2)
    add_table(doc, ["Code", "Screen/Function"], [
        ["B21", "Portfolio summary"], ["B22", "Account details"], ["B23", "Account statement"],
        ["B24", "Loan repayment"], ["B25", "Recent transactions"], ["B26", "Interest certificate"],
        ["B27", "Disbursement schedule"], ["B28", "New loan request"], ["B29", "Disbursement request"],
        ["B2A", "Reschedule request"], ["B2B", "Top-up request"], ["B2C", "Closure simulation"],
        ["B1", "Multi-factor authentication"], ["B2", "User profiling"], ["B3", "Multi-level entitlements"],
        ["B4", "Personalisation"], ["B5", "Workflow"], ["B6", "Alerts & notifications"],
    ])

    add_heading(doc, "6.3 K-codes — Reports & Interfaces", 2)
    add_table(doc, ["Code", "Function"], [
        ["K211", "Standard reports"], ["K212", "Accounting entries query"],
        ["K213", "Pre-closure simulation"], ["K214", "Compliance reporting"],
        ["K215", "Central bank reporting"], ["K216", "Repayment reports"],
        ["K217", "Approval & disbursal reports"], ["K218", "Collateral reports"],
        ["K219", "Operational analytics"], ["K21A", "Accrual reports"],
        ["K21B", "Breach reports (collateral & limits)"], ["K21C", "Delinquency reports"],
        ["K221", "Accounting entries → GL"], ["K222", "Credit score interface"],
        ["K223", "Rates / FX interface"], ["K224", "Multi-channel payments"],
        ["K225", "Customer / accounts data"], ["K226", "KYC & blacklist"],
    ])

    # Section 7
    add_heading(doc, "7. Product Configuration", 1)
    add_para(doc, "Policy layers attached to each product:", bold=True)
    add_table(doc, ["Policy", "Controls"], [
        ["Interest policy", "Rate, margin, reset frequency"],
        ["Fee policy", "Processing, maintenance fees"],
        ["Priority policy", "Payment allocation order"],
        ["Delinquency policy", "Bucket rules, NPA triggers"],
        ["Accounting policy", "GL accounts, accrual method"],
        ["Penalty charges policy", "Late payment penalties"],
    ])

    add_para(doc, "Personal Loan Setup Order:", bold=True)
    for i, step in enumerate([
        "Create base product (Retail → Unsecured → Fixed → EMI)",
        "Define rate scheme (e.g. 12% fixed)",
        "Define fee scheme (e.g. 2% processing)",
        "Define priority scheme (Fees → Interest → Principal)",
        "Define delinquency scheme (30/60/90 DPD buckets)",
        "Define accounting scheme (GL codes)",
        "Create specific product (min/max amount, tenor, prepayment rules)",
        "Test: simulator → booking → disbursal → payment → EOD",
    ], 1):
        add_bullet(doc, f"{i}. {step}")

    # Section 8
    add_heading(doc, "8. Daily Transactions vs EOD Batch", 1)
    add_code_block(doc, """DAY (Online / Branch / API)
  • Disbursal, Scheduled Payment, Reschedule, Waiver
  • Immediate validation + partial accounting

NIGHT (G28 Batch / EOD)
  • Interest accrual for all active loans
  • Due creation for due dates
  • Delinquency stamping (DPD update)
  • Late fees, provisions, consolidated GL postings

Payment Allocation Example (Priority Scheme):
  Payment ₹10,000 received
    → ₹500  Late fee
    → ₹200  Other charges
    → ₹3,000 Interest due
    → ₹6,300 Principal""")

    # Section 9
    add_heading(doc, "9. Integration Layer", 1)
    add_table(doc, ["Integration", "Typical Flow"], [
        ["LOS → LMS booking", "API event on approval"],
        ["LMS → Core GL", "K221 accounting entries"],
        ["LMS → Payments", "K224 multi-channel payments"],
        ["Bureau", "K222 during origination"],
        ["KYC", "K226 blacklist check"],
        ["DEP → LMS", "B24 repayment API"],
        ["Olive Fabric", "Custom orchestration flows"],
    ])

    add_code_block(doc, """External: Tax DB · ROC · Credit Bureau · E-Sign · Bank Statement Analyzer
                    ↓
              iTurmeric (API + Events)
                    ↓
         LOS · LMS · CLMS · Core · CRM · DWH · ECM · SMS/Email""")

    # Section 10
    add_heading(doc, "10. Server & Environment Access", 1)
    add_para(doc, "Ask your trainer for this access checklist:", bold=True)
    add_table(doc, ["Item", "What to request"], [
        ["LMS UI URL", "Training/SIT environment URL"],
        ["LOS URL", "Same or separate portal with role-based menu"],
        ["Weaver/iTurmeric", "Integration studio URL"],
        ["Environment", "DEV / SIT / UAT / TRAINING"],
        ["Credentials", "Trainee ID + role (Ops, Product Admin, Developer)"],
        ["VPN", "Required for internal IPs (e.g. 172.16.x.x)"],
        ["Database", "Read-only access (Oracle/PostgreSQL)"],
        ["API gateway", "Swagger/OpenAPI docs for LMS v22 APIs"],
        ["Batch monitor", "G28 EOD job status screen"],
        ["Log location", "Application logs for failed booking/payment"],
    ])

    add_para(doc, "Typical Roles:", bold=True)
    add_table(doc, ["Role", "Can access"], [
        ["Product Admin", "G21 product setup"],
        ["Loan Ops", "G25 booking, G26 transactions"],
        ["Supervisor", "Waivers, reversals, manual DPD"],
        ["Batch Admin", "G28 EOD monitoring"],
        ["Developer", "iTurmeric, API config"],
        ["Read-only", "Queries, reports (K21x)"],
    ])

    # Section 11
    add_heading(doc, "11. Study Checklist & Screenshots", 1)
    add_para(doc, "Capture these screenshots during training:", bold=True)
    add_table(doc, ["#", "Screenshot Topic", "Why"], [
        ["1", "LMS main menu (G21–G28)", "Module map"],
        ["2", "Product Definition screen", "Product engine"],
        ["3", "Rate / Fee / Delinquency schemes", "Policy setup"],
        ["4", "Retail Loan Booking (G25)", "LOS handoff"],
        ["5", "Disbursal screen", "Fund release"],
        ["6", "Repayment schedule after booking", "EMI structure"],
        ["7", "Scheduled Payment (G26)", "Day transaction"],
        ["8", "Payment allocation breakdown", "Priority scheme"],
        ["9", "Reschedule / Deferral screen", "Modifications"],
        ["10", "Delinquency bucket view (G27)", "DPD"],
        ["11", "EOD batch monitor (G28)", "Batch jobs"],
        ["12", "GL/accounting entries (K212)", "Finance link"],
        ["13", "CLMS collateral link (G24)", "Security"],
        ["14", "Limit utilization (G23)", "Exposure"],
        ["15", "DEP B24 repayment journey", "Channel integration"],
    ])

    add_para(doc, "Topics to Master:", bold=True)
    for topic in [
        "Difference between LOS and LMS",
        "What happens in EOD vs online transaction",
        "Payment priority allocation order",
        "Fixed vs floating rate reset",
        "Single vs multi-tranche disbursal",
        "DPD buckets and NPA classification",
        "Top-up vs rollover vs reschedule",
        "CLMS collateral lifecycle (earmark → release)",
        "LOS → LMS handoff data flow",
        "G25 vs G26 vs G28",
    ]:
        add_bullet(doc, f"☐ {topic}")

    # Section 12
    add_heading(doc, "12. Glossary", 1)
    add_table(doc, ["Term", "Definition"], [
        ["Booking", "Create live loan account in LMS after approval"],
        ["Disbursal", "Actual fund transfer to customer"],
        ["Accrual", "Recognize interest income/expense over time"],
        ["Due", "Installment amount owed on a date"],
        ["DPD", "Days Past Due"],
        ["NPA", "Non-Performing Asset"],
        ["STP", "Straight-Through Processing (no manual step)"],
        ["Moratorium / Deferral", "Payment holiday"],
        ["POS", "Principal Outstanding"],
        ["GL", "General Ledger"],
        ["CAM", "Credit Appraisal Memo (corporate)"],
        ["Drawing power", "Amount customer can draw against limit/collateral"],
        ["PF Credit", "Purple Fabric AI experts for lending"],
    ])

    # Section 13
    add_heading(doc, "13. Official References (Public)", 1)
    refs = [
        "Retail Loan Management: https://www.intellectdesign.com/solutions/lending-solutions/retail-loan-management/",
        "Retail Loan Originations: https://www.intellectdesign.com/solutions/lending-solutions/retail-loan-originations/",
        "Commercial Loan Management: https://www.intellectdesign.com/solutions/lending-solutions/commercial-loan-management/",
        "Commercial Loan Originations: https://www.intellectdesign.com/solutions/lending-solutions/commercial-loan-originations/",
        "Lending Brochure PDF (G/B/K codes): https://www.intellectdesign.com/wp-content/uploads/2025/12/lending-brochure-spread.pdf",
        "eMACH.ai Technology: https://www.intellectdesign.com/technology/",
        "iTurmeric: https://www.intellectdesign.com/iturmeric/",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    return doc


if __name__ == "__main__":
    output_path = r"c:\Users\kirub\Documents\global_logistics_second\docs\Intellect_LMS_v22_Study_Guide.docx"
    document = build_document()
    document.save(output_path)
    print(f"Saved: {output_path}")
